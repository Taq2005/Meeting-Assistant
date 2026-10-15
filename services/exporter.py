"""Export a MeetingResult to DOCX or PDF."""
from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                Table, TableStyle)

from core.models import MeetingResult

BLUE_HEX = "#1E5BD8"
LIGHT_HEX = "#EAF1FD"
BLUE_RGB = RGBColor(0x1E, 0x5B, 0xD8)
NOT_SPECIFIED = "Not specified."


def _doc_title(result: MeetingResult) -> str:
    return result.context.title.strip() or "Meeting Report"


def _date_line(result: MeetingResult) -> str:
    return result.context.date.strip()


# ---------------------------------------------------------------- DOCX
def _docx_heading(doc, text: str, level: int = 1) -> None:
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = BLUE_RGB


def _docx_bullets(doc, items) -> None:
    if not items:
        doc.add_paragraph(NOT_SPECIFIED)
        return
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def to_docx(result: MeetingResult) -> bytes:
    doc = Document()
    _docx_heading(doc, _doc_title(result), level=0)
    if _date_line(result):
        doc.add_paragraph(f"Date: {_date_line(result)}")

    _docx_heading(doc, "Executive Summary")
    doc.add_paragraph(result.summary or NOT_SPECIFIED)

    m = result.minutes
    _docx_heading(doc, "Minutes of Meeting")
    _docx_heading(doc, "Attendees", level=2)
    _docx_bullets(doc, m.attendees)
    _docx_heading(doc, "Agenda Items", level=2)
    _docx_bullets(doc, m.agenda_items)
    _docx_heading(doc, "Discussion", level=2)
    if m.discussion:
        for para in m.discussion:
            doc.add_paragraph(para)
    else:
        doc.add_paragraph(NOT_SPECIFIED)
    _docx_heading(doc, "Decisions", level=2)
    _docx_bullets(doc, m.decisions)
    _docx_heading(doc, "Action Items", level=2)
    if m.action_items:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        for cell, label in zip(table.rows[0].cells, ("Task", "Owner", "Due")):
            cell.text = label
        for item in m.action_items:
            row = table.add_row().cells
            row[0].text, row[1].text, row[2].text = item.task, item.owner, item.due
    else:
        doc.add_paragraph(NOT_SPECIFIED)

    _docx_heading(doc, "Key Points")
    _docx_bullets(doc, result.key_points)
    _docx_heading(doc, "Recommendations")
    _docx_bullets(doc, result.recommendations)

    doc.add_page_break()
    _docx_heading(doc, "Transcript")
    for line in result.transcript.splitlines():
        if line.strip():
            doc.add_paragraph(line)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ----------------------------------------------------------------- PDF
def _pdf_styles():
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("T", parent=base["Title"], textColor=colors.HexColor(BLUE_HEX)),
        "h1": ParagraphStyle("H1", parent=base["Heading1"], textColor=colors.HexColor(BLUE_HEX)),
        "h2": ParagraphStyle("H2", parent=base["Heading2"], textColor=colors.HexColor(BLUE_HEX)),
        "body": ParagraphStyle("B", parent=base["BodyText"], leading=15),
        "cell": ParagraphStyle("C", parent=base["BodyText"], fontSize=9, leading=12),
        "head": ParagraphStyle("Hd", parent=base["BodyText"], fontSize=9, leading=12,
                               textColor=colors.white),
    }


def _pdf_bullets(story, items, styles) -> None:
    if not items:
        story.append(Paragraph(NOT_SPECIFIED, styles["body"]))
        return
    for item in items:
        story.append(Paragraph(escape(item), styles["body"], bulletText="•"))


def to_pdf(result: MeetingResult) -> bytes:
    styles = _pdf_styles()
    story = [Paragraph(escape(_doc_title(result)), styles["title"])]
    if _date_line(result):
        story.append(Paragraph(f"Date: {escape(_date_line(result))}", styles["body"]))
    story.append(Spacer(1, 0.3 * cm))

    story.append(Paragraph("Executive Summary", styles["h1"]))
    story.append(Paragraph(escape(result.summary or NOT_SPECIFIED), styles["body"]))

    m = result.minutes
    story.append(Paragraph("Minutes of Meeting", styles["h1"]))
    story.append(Paragraph("Attendees", styles["h2"]))
    _pdf_bullets(story, m.attendees, styles)
    story.append(Paragraph("Agenda Items", styles["h2"]))
    _pdf_bullets(story, m.agenda_items, styles)
    story.append(Paragraph("Discussion", styles["h2"]))
    if m.discussion:
        for para in m.discussion:
            story.append(Paragraph(escape(para), styles["body"]))
            story.append(Spacer(1, 0.15 * cm))
    else:
        story.append(Paragraph(NOT_SPECIFIED, styles["body"]))
    story.append(Paragraph("Decisions", styles["h2"]))
    _pdf_bullets(story, m.decisions, styles)
    story.append(Paragraph("Action Items", styles["h2"]))
    if m.action_items:
        rows = [[Paragraph(label, styles["head"]) for label in ("Task", "Owner", "Due")]]
        for item in m.action_items:
            rows.append([Paragraph(escape(v), styles["cell"])
                         for v in (item.task, item.owner, item.due)])
        table = Table(rows, colWidths=[9 * cm, 4 * cm, 3.5 * cm], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BLUE_HEX)),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(LIGHT_HEX)]),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8CCF2")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(table)
    else:
        story.append(Paragraph(NOT_SPECIFIED, styles["body"]))

    story.append(Paragraph("Key Points", styles["h1"]))
    _pdf_bullets(story, result.key_points, styles)
    story.append(Paragraph("Recommendations", styles["h1"]))
    _pdf_bullets(story, result.recommendations, styles)

    story.append(PageBreak())
    story.append(Paragraph("Transcript", styles["h1"]))
    for line in result.transcript.splitlines():
        if line.strip():
            story.append(Paragraph(escape(line), styles["cell"]))

    buffer = BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                      topMargin=2 * cm, bottomMargin=2 * cm,
                      title=_doc_title(result)).build(story)
    return buffer.getvalue()
